"""用途：验证分类版和层级版 Excel 导出脚本，不是实际批次转换入口。

运行方法（在 PDFToTex 根目录执行）：
    python -m pytest -q scripts/batch_excel/test_exporters.py
只运行某类用例：
    python -m pytest -q scripts/batch_excel/test_exporters.py -k hierarchy

依赖：Python 3.10+、pytest、openpyxl，可用 python -m pip install pytest openpyxl 安装。
无需填写真实批次或配置输入输出路径；测试自行构造 JSON，使用 pytest 临时目录写入 Excel。
不会读取或覆盖 data/history 中的正式结果，也不会调用模型接口或启动 PDF 流水线。
主要覆盖：任意批次、新属性列、角色差异、表单边界、空对象、编号保真、
覆盖保护、参数/顶部路径配置、非法 JSON 拒绝，以及导出文件内容和无配色要求。
退出码 0 表示通过；其他退出码应查看终端显示的失败用例或环境错误。
"""

import importlib.util
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

HERE = Path(__file__).resolve().parent


def fixture():
    return {'batch': 'B-SECOND', 'process': [{'name': 'process one', 'subprocess': [
        {'name': 'phase', 'step': [{'name': 'step', 'form': [
            {'name': 'same form', 'personnel': [
                {'name': 'Person', 'attributes': {'role': 'operator', 'date': '2025-09-02'}},
                {'name': 'Person', 'attributes': {'date': '2025-09-02', 'role': 'operator'}},
                {'name': 'Person', 'attributes': {'role': 'reviewer'}},
            ], 'materials': [{'name': 'Reagent', 'attributes': {
                'lot_number': '00123', 'quantity': 0, 'approved': False,
                'notes': '=1+1', 'name': 'secondary',
                'code': 12345678901234567, 'brand': None,
            }}], 'equipment': [{'name': 'Device', 'attributes': {}}], 'environment': []},
            {'name': 'same form', 'personnel': [
                {'name': 'Person', 'attributes': {'role': 'operator', 'date': '2025-09-02'}}
            ], 'materials': [], 'equipment': [], 'environment': []},
            {'name': 'empty form', 'personnel': [], 'materials': [], 'equipment': [], 'environment': []},
        ]}]}]}]}


def run(tmp_path, script, data, *args):
    source = tmp_path / 'input.json'
    source.write_text(json.dumps(data), encoding='utf-8')
    result = subprocess.run([sys.executable, str(HERE / script), str(source),
                             '--output-dir', str(tmp_path / 'out'), *args],
                            cwd=tmp_path, capture_output=True, text=True)
    return result, source


def expand(ws):
    rows = [list(row) for row in ws.values]
    for region in ws.merged_cells.ranges:
        for r in range(region.min_row, region.max_row + 1):
            rows[r - 1][region.min_col - 1] = ws.cell(region.min_row, region.min_col).value
    return rows


@pytest.mark.parametrize('script,suffix', [('build_workbook.py', 'categories'), ('build_hierarchy.py', 'hierarchy')])
def test_cli_uses_input_batch_and_does_not_change_json(tmp_path, script, suffix):
    data = fixture()
    result, source = run(tmp_path, script, data)
    assert result.returncode == 0, result.stderr
    assert json.loads(source.read_text()) == data
    book = load_workbook(tmp_path / 'out' / f'B-SECOND_{suffix}.xlsx')
    assert all(cell.data_type != 'f' and not cell.font.bold and cell.fill.patternType is None
               for ws in book for row in ws for cell in row)
    assert all('form_order' not in [cell.value for cell in ws[1]] for ws in book)
    assert all('批次' in next(ws.values) for ws in book)
    assert all(ws.freeze_panes == 'A2' for ws in book)


def test_categories_preserve_dynamic_attributes_roles_and_cross_form_occurrences(tmp_path):
    result, _ = run(tmp_path, 'build_workbook.py', fixture())
    assert result.returncode == 0, result.stderr
    book = load_workbook(tmp_path / 'out/B-SECOND_categories.xlsx')
    assert book.sheetnames == ['步骤总览', '人员明细', '物料明细', '设备明细', '环境明细']
    overview = list(book['步骤总览'].values)
    assert len(overview) == 4
    assert [r[4] for r in overview[1:]] == ['same form', 'same form', 'empty form']
    assert overview[0][-4:] == ('人员记录数', '物料记录数', '设备记录数', '环境记录数')
    assert [r[-4:] for r in overview[1:]] == [(2, 1, 1, 0), (1, 0, 0, 0), (0, 0, 0, 0)]
    people = list(book['人员明细'].values)
    assert len(people) == 4
    assert [r[people[0].index('角色')] for r in people[1:]] == ['operator', 'reviewer', 'operator']
    material = list(book['物料明细'].values)
    row = dict(zip(material[0], material[1]))
    assert row['物料名称'] == 'Reagent' and row['名称'] == 'secondary'
    assert row['备注'] == '=1+1' and row['批号'] == '00123'
    assert row['数量'] == 0 and row['是否批准'] is False
    assert row['代码'] == '12345678901234567'
    assert book['设备明细'].max_row == 2
    assert book['环境明细'].max_row == 1


def test_hierarchy_preserves_empty_objects_and_form_boundaries(tmp_path):
    result, _ = run(tmp_path, 'build_hierarchy.py', fixture())
    assert result.returncode == 0, result.stderr
    ws = load_workbook(tmp_path / 'out/B-SECOND_hierarchy.xlsx').active
    rows = expand(ws)
    assert ws.title == '步骤聚合层级表'
    assert rows[0] == ['批次', '大工序', '阶段', '步骤', '步骤/表单', '类别', '对象名称', '属性', '值']
    assert [r[-1] for r in rows if r[-2] == '角色'] == ['operator', 'reviewer', 'operator']
    assert {r[5] for r in rows[1:]} == {'人员', '物料', '设备', None}
    assert any(r[6] == 'Device' and r[7:] == [None, None] for r in rows[1:])
    assert rows[-1][4] == 'empty form' and rows[-1][5:] == [None] * 4
    forms = [region for region in ws.merged_cells.ranges if region.min_col == 5]
    assert len(forms) == 2
    assert forms[0].max_row < forms[1].min_row or forms[1].max_row < forms[0].min_row


@pytest.mark.parametrize('script', ['build_workbook.py', 'build_hierarchy.py'])
def test_existing_output_requires_overwrite(tmp_path, script):
    result, _ = run(tmp_path, script, fixture())
    assert result.returncode == 0, result.stderr
    target = next((tmp_path / 'out').glob('*.xlsx'))
    before = target.read_bytes()
    result, _ = run(tmp_path, script, fixture())
    assert result.returncode != 0
    assert target.read_bytes() == before
    result, _ = run(tmp_path, script, fixture(), '--overwrite')
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize('bad', ['missing', 'nested', 'chinese', 'order', 'control', 'long'])
def test_invalid_schema_fails_without_creating_output(tmp_path, bad):
    data = fixture()
    attrs = data['process'][0]['subprocess'][0]['step'][0]['form'][0]['personnel'][0]['attributes']
    if bad == 'missing':
        del data['process'][0]['subprocess']
    else:
        key, value = {
            'nested': ('extra', {'nested': 'not scalar'}), 'chinese': ('角色', 'operator'),
            'order': ('form_order', 1), 'control': ('extra', '\x01'), 'long': ('extra', 'a' * 32768),
        }[bad]
        attrs[key] = value
    result, _ = run(tmp_path, 'build_hierarchy.py', data)
    assert result.returncode != 0
    assert 'error:' in result.stderr.lower()
    assert not list(tmp_path.rglob('*.xlsx'))


@pytest.mark.parametrize('script', ['build_workbook.py', 'build_hierarchy.py'])
def test_empty_batch_and_configuration_api(tmp_path, script):
    result, _ = run(tmp_path, script, {'batch': 'EMPTY', 'process': []})
    assert result.returncode == 0, result.stderr
    # Importing does not export files or require access to a real batch.
    sys.path.insert(0, str(HERE))
    spec = importlib.util.spec_from_file_location(script[:-3], HERE / script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    source = tmp_path / 'configured.json'
    source.write_text(json.dumps({'batch': 'CONFIG', 'process': []}))
    module.INPUT_JSON = source
    module.OUTPUT_DIR = tmp_path / 'configured'
    module.main([])
    assert len(list((tmp_path / 'configured').glob('CONFIG_*.xlsx'))) == 1


@pytest.mark.parametrize('script,suffix', [('build_workbook.py', 'categories'), ('build_hierarchy.py', 'hierarchy')])
def test_default_output_directory_and_duplicate_json_key_rejection(tmp_path, script, suffix):
    source = tmp_path / 'input.json'
    source.write_text('{"batch":"DEFAULT","process":[]}')
    result = subprocess.run([sys.executable, str(HERE / script), str(source)],
                            cwd=tmp_path.parent, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    target = tmp_path / f'DEFAULT_{suffix}.xlsx'
    assert target.exists()
    before = target.read_bytes()
    source.write_text('{"batch":"DEFAULT","batch":"OTHER","process":[]}')
    result = subprocess.run([sys.executable, str(HERE / script), str(source), '--overwrite'],
                            capture_output=True, text=True)
    assert result.returncode != 0 and 'Duplicate JSON key' in result.stderr
    assert target.read_bytes() == before


@pytest.mark.parametrize('script,suffix', [('build_workbook.py', 'categories'), ('build_hierarchy.py', 'hierarchy')])
def test_timestamp_display_keeps_time_including_midnight(tmp_path, script, suffix):
    import datetime as dt
    data = fixture()
    obj = data['process'][0]['subprocess'][0]['step'][0]['form'][0]['materials'][0]
    obj['attributes'] = {'production_date': '2026-09-15', 'start_time': '2026-09-15-16:00:05',
                         'end_time': '2026-09-16-00:00:00', 'sampling_time': '2026-09-15-16:01',
                         'time': '16:02'}
    result, _ = run(tmp_path, script, data)
    assert result.returncode == 0, result.stderr
    wb = load_workbook(tmp_path / 'out' / f'B-SECOND_{suffix}.xlsx')
    dates = [cell for ws in wb for row in ws for cell in row if isinstance(cell.value, dt.datetime)]
    afternoon = [cell for cell in dates if cell.value == dt.datetime(2026, 9, 15, 16, 0, 5)]
    midnight = [cell for cell in dates if cell.value == dt.datetime(2026, 9, 16)]
    assert len(afternoon) == len(midnight) == 1
    assert afternoon[0].number_format == midnight[0].number_format == 'yyyy-mm-dd-hh:mm:ss'
    minute = [cell for cell in dates if cell.value == dt.datetime(2026, 9, 15, 16, 1)]
    assert len(minute) == 1 and minute[0].number_format == 'yyyy-mm-dd-hh:mm'
    clocks = [cell for ws in wb for row in ws for cell in row if cell.value == dt.time(16, 2)]
    assert len(clocks) == 1 and clocks[0].number_format == 'hh:mm'


@pytest.mark.parametrize('script,suffix', [('build_workbook.py', 'categories'), ('build_hierarchy.py', 'hierarchy')])
def test_flat_schema_keeps_step_and_object_ids(tmp_path, script, suffix):
    data = {'batch': 'NEW', 'process': [{'subprocess': 'Release', 'step': [
        {'id': 'S001', 'stage': '', 'form': 'Form', 'personnel': [
            {'id': 'O001', 'name': 'Person', 'attributes': {'role': 'operator'}},
            {'id': 'O002', 'name': 'Person', 'attributes': {'role': 'reviewer'}},
        ], 'materials': [], 'equipment': [], 'environment': []},
        {'id': 'S002', 'stage': '', 'form': 'Form', 'personnel': [
            {'id': 'O003', 'name': 'Person', 'attributes': {'role': 'operator'}},
        ], 'materials': [], 'equipment': [], 'environment': []}
    ]}]}
    result, _ = run(tmp_path, script, data)
    assert result.returncode == 0, result.stderr
    wb = load_workbook(tmp_path / 'out' / f'NEW_{suffix}.xlsx')
    ws = wb['人员明细'] if suffix == 'categories' else wb.active
    rows = expand(ws)
    si, oi = rows[0].index('步骤ID'), rows[0].index('记录ID')
    assert [(r[si], r[oi]) for r in rows[1:]] == [('S001', 'O001'), ('S001', 'O002'), ('S002', 'O003')]
    assert '大工序' in rows[0] and '表单序次' not in rows[0]
    if suffix == 'categories':
        assert rows[0][:7] == ['记录ID', '步骤ID', '批次', '大工序', '阶段', '步骤/表单', '人员名称']
        assert list(wb['步骤总览'].values)[1:] == [
            ('S001', 'NEW', 'Release', None, 'Form', 2, 0, 0, 0),
            ('S002', 'NEW', 'Release', None, 'Form', 1, 0, 0, 0)]


@pytest.mark.parametrize('script', ['build_workbook.py', 'build_hierarchy.py'])
def test_unmapped_attribute_rejected_without_replacing_output(tmp_path, script):
    data = fixture()
    result, _ = run(tmp_path, script, data)
    assert result.returncode == 0, result.stderr
    target = next((tmp_path / 'out').glob('*.xlsx'))
    before = target.read_bytes()
    data['process'][0]['subprocess'][0]['step'][0]['form'][0]['materials'][0]['attributes']['unmapped_field'] = 'value'
    result, _ = run(tmp_path, script, data, '--overwrite')
    assert result.returncode != 0
    assert 'unmapped_field' in result.stderr and 'labels_zh.py' in result.stderr
    assert target.read_bytes() == before


def test_translated_column_collision_preserves_both_values(tmp_path):
    data = fixture()
    data['process'][0]['subprocess'][0]['step'][0]['form'][0]['equipment'][0]['attributes'] = {'equipment_name': 'Other name'}
    result, _ = run(tmp_path, 'build_workbook.py', data)
    assert result.returncode == 0, result.stderr
    wb = load_workbook(tmp_path / 'out/B-SECOND_categories.xlsx')
    rows = list(wb['设备明细'].values)
    obj = dict(zip(rows[0], rows[1]))
    assert obj['设备名称'] == 'Device'
    assert obj['属性：设备名称'] == 'Other name'
