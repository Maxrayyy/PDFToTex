"""用途：将约定的英文批次 JSON 导出为分类 Excel，每个对象占一行。

输出工作表：步骤总览、人员明细、物料明细、设备明细、环境明细，均使用中文表头。
依赖：Python 3.10+、openpyxl；复制时需保留 common.py、temporal.py、labels_zh.py。
安装依赖（在本脚本目录执行）：python -m pip install -r requirements.txt

命令行用法（以下示例在本脚本目录执行；其他目录请使用脚本的完整路径）：
    python build_workbook.py /path/to/batch.json
    python build_workbook.py /path/to/batch.json --output-dir /path/to/output
    python build_workbook.py /path/to/batch.json --output-dir /path/to/output --overwrite

路径配置：也可修改下方 INPUT_JSON 和 OUTPUT_DIR，然后直接运行 python build_workbook.py。
建议填写绝对路径；相对路径以运行命令时的当前目录为准，而不是脚本所在目录。
命令行参数优先于脚本配置；OUTPUT_DIR 为 None 时，输出到输入 JSON 所在目录。
输出文件：<batch>_categories.xlsx，其中 batch 从 JSON 中读取，不需要另填批次号。
默认不覆盖已有文件；确认要替换时添加 --overwrite。不会修改输入 JSON。

输入结构：batch -> process[{subprocess, step: [{id, stage, form, 四类对象数组}]}]。
subprocess 是工序名称，form 是表单名称字符串；对象为 {id, name, attributes}。
兼容旧版各层 name 包装的数组结构；完整 JSON 示例见同目录 README.md。
日期显示为 YYYY-MM-DD，带时间时用连字符连接；只有时分时不补秒。
attributes 使用英文 snake_case 字段，值为文本、数字、布尔值或 null，不支持嵌套数组或对象。
新增属性先在 labels_zh.py 添加中文映射；缺失映射会报错，不会覆盖已有结果。
属性中文名与已有表头重名时增加“属性：”前缀，防止同名列混淆。
仅合并同一表单、同一类别内名称和全部属性完全一致的对象；角色或日期不同的不会合并。
不添加表单次序、来源列、颜色或装饰样式。写入后回读核对成功才发布最终文件。
仅冻结首行表头，不冻结左侧列；横向滚动时所有列一起移动。
步骤总览包含四类记录数，为本次导出的去重后明细记录数量；空类别显示 0。
记录数不是人数或独立物料种类数，手工修改 Excel 明细后需重新导出以更新统计。
"""

INPUT_JSON = None  # 输入 JSON 的绝对路径，例如 "/absolute/path/to/batch.json"。
OUTPUT_DIR = None  # 输出目录的绝对路径；None 表示使用输入 JSON 所在目录。

from openpyxl import Workbook

from common import CATEGORIES, PREFIX, append, cli, excel_value, flat_schema, forms, layout
from labels_zh import ATTRIBUTE_LABELS, CATEGORY_LABELS, require_labels, structural_label


def build(data):
    flat = flat_schema(data)
    records = list(forms(data))
    require_labels(records)
    groups = {category: [] for category in CATEGORIES}
    for _, prefix, form in records:
        for category in CATEGORIES:
            for obj in form[category]:
                groups[category].append((prefix + ([obj['id']] if flat else []) + [obj['name']], obj['attributes']))
    wb = Workbook()
    wb.remove(wb.active)
    expected = {}
    ws = wb.create_sheet('步骤总览')
    overview_keys = ['step_id', 'batch', 'subprocess', 'stage', 'form'] if flat else list(PREFIX[:-1])
    overview = [[structural_label(key, flat) for key in overview_keys]
                + [CATEGORY_LABELS[category] + '记录数' for category in CATEGORIES]]
    overview.extend(([prefix[3], *prefix[:3], prefix[4]] if flat else prefix)
                    + [len(categories[category]) for category in CATEGORIES]
                    for _, prefix, categories in records)
    for row in overview:
        append(ws, row)
    layout(ws, ([16, 25, 42, 34, 58] if flat else [25, 42, 34, 34, 58]) + [18] * 4)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    expected[ws.title] = overview
    for category, objects in groups.items():
        fields = list(dict.fromkeys(key for _, attrs in objects for key in attrs))
        headers = ['batch', 'subprocess', 'stage', 'step_id', 'form', 'id', 'name'] if flat else list(PREFIX)
        order = [5, 3, 0, 1, 2, 4, 6] if flat else list(range(len(headers)))
        headers = [CATEGORY_LABELS[category] + '名称' if headers[i] == 'name'
                   else structural_label(headers[i], flat) for i in order]
        prefix_size = len(headers)
        for key in fields:
            label = ATTRIBUTE_LABELS[key]
            while label in headers:
                label = '属性：' + label
            headers.append(label)
        ws = wb.create_sheet(CATEGORY_LABELS[category] + '明细')
        expected[ws.title] = [headers]
        append(ws, headers)
        for prefix, attrs in objects:
            row = [prefix[i] for i in order] + [excel_value(key, attrs.get(key)) for key in fields]
            append(ws, row)
            expected[ws.title].append(row)
        widths = ([16, 16, 25, 42, 34, 38, 48] if flat else [25, 42, 34, 34, 38, 48]) + [min(60, max(26, len(h) * 2 + 4)) for h in headers[prefix_size:]]
        layout(ws, widths)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    return wb, expected


def main(argv=None):
    return cli(build, 'categories', argv, INPUT_JSON, OUTPUT_DIR)


if __name__ == '__main__':
    main()
